# -*- coding: utf-8 -*-
"""Rebuild CHUONG III Word document from project-aligned content."""
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT_PATH = r"e:\UNUtrip\docx_doan\chuong_3_demo_final_v3.docx"
ALT_PATH = r"e:\UNUtrip\docx_doan\chuong_3_demo_final_v3_rebuilt.docx"


def add_para(doc, text, bold=False, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(12)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(text, style="List Bullet")
    for run in p.runs:
        run.font.size = Pt(12)


def set_doc_defaults(doc):
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)
    try:
        style._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}ascii", "Times New Roman")
        style._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}hAnsi", "Times New Roman")
        style._element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}cs", "Times New Roman")
    except Exception:
        pass


def build():
    doc = Document()
    set_doc_defaults(doc)

    # --- Title ---
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("CHƯƠNG III\nTHIẾT KẾ HỆ THỐNG")
    r.bold = True
    r.font.size = Pt(14)

    doc.add_paragraph()

    intro = (
        "Chương này trình bày thiết kế kỹ thuật của hệ thống ứng dụng du lịch thông minh "
        "UNU Trip trên nền tảng Android, tích hợp backend REST, cơ sở dữ liệu quan hệ và "
        "phân hệ truy vấn–sinh câu trả lời dựa trên tri thức (RAG). Nội dung được căn cứ "
        "vào kiến trúc và mã nguồn hiện có của dự án, nhằm đảm bảo tính nhất quán giữa "
        "tài liệu và triển khai thực tế."
    )
    add_para(doc, intro)

    # ========== 3.1 ==========
    doc.add_heading("3.1. Tổng quan thiết kế hệ thống", level=1)

    doc.add_heading("3.1.1. Mục tiêu thiết kế", level=2)
    add_para(
        doc,
        "Mục tiêu thiết kế của hệ thống là xây dựng một kiến trúc phân tán, dễ mở rộng và "
        "bảo trì, trong đó ứng dụng di động đảm nhiệm trải nghiệm người dùng; máy chủ ứng dụng "
        "đảm nhiệm nghiệp vụ, xác thực và truy cập dữ liệu; dịch vụ RAG đảm nhiệm truy hồi "
        "thông tin và hỗ trợ sinh nội dung có kiểm soát cho các luồng AI. Đồng thời, thiết kế "
        "ưu tiên tách biệt trách nhiệm theo từng phân hệ (mobile, API, dữ liệu, AI) để giảm "
        "độ phụ thuộc và thuận tiện kiểm thử theo từng lớp.",
    )

    doc.add_heading("3.1.2. Nguyên tắc thiết kế", level=2)
    bullets = [
        "Tách lớp presentation / domain / data trên client Android (Fragment–ViewModel–Repository–API).",
        "Backend Node.js tổ chức theo module nghiệp vụ (routes → controller → service → repository).",
        "Giao tiếp giữa các thành phần chủ yếu qua HTTP và JSON; định danh phiên làm việc người dùng bằng JWT.",
        "Dữ liệu nghiệp vụ lưu trữ tập trung trên MySQL; phân hệ RAG có pipeline và endpoint riêng, được backend proxy có kiểm soát.",
        "Cấu hình môi trường (URL API, URL RAG, thông tin DB, khóa nội bộ) tách khỏi mã nguồn thông qua biến môi trường và build config.",
    ]
    for b in bullets:
        add_bullet(doc, b)

    doc.add_heading("3.1.3. Kiến trúc tổng thể hệ thống", level=2)
    add_para(
        doc,
        "Theo triển khai hiện tại, hệ thống gồm các lớp logic sau: (i) ứng dụng Android UNU Trip "
        "(applicationId sản phẩm com.smarttravel; biến thể dev có hậu tố .dev); (ii) máy chủ "
        "Node.js/Express phục vụ REST dưới tiền tố /api, phục vụ tài nguyên tĩnh uploads/images "
        "và dashboard quản trị web /admin; (iii) MySQL (mặc định cơ sở dữ liệu unudata) lưu người "
        "dùng, địa điểm (bảng app_places), yêu thích, đánh giá, lịch trình; (iv) dịch vụ FastAPI "
        "RAG (cổng publish mặc định 8001 trong docker-compose) xử lý truy vấn BM25/hybrid và các "
        "endpoint chat/itinerary phía AI; (v) Redis trong stack Docker hỗ trợ vận hành RAG (rate "
        "limit/cache theo cấu hình).",
    )
    add_para(
        doc,
        "(Hình 3.1 — Sơ đồ tổng quan hệ thống: chèn ảnh UNU_Trip_So_do_tong_quan_he_thong.png hoặc "
        "UNU_Trip_Kien_truc_Client_Server_AI_RAG.png tại thư mục gốc dự án.)",
        italic=True,
    )

    doc.add_heading("3.1.4. Các phân hệ chính của hệ thống", level=2)
    tbl = doc.add_table(rows=1, cols=4)
    hdr = tbl.rows[0].cells
    hdr[0].text = "Phân hệ"
    hdr[1].text = "Công nghệ"
    hdr[2].text = "Vai trò"
    hdr[3].text = "Ghi chú triển khai"
    rows = [
        ("Client Android", "Kotlin, Navigation, Retrofit, ViewBinding", "Đăng nhập, duyệt địa điểm, lịch trình, chatbot, bản đồ OSMDroid", "BASE_URL qua BuildConfig; thời tiết gọi trực tiếp Open-Meteo"),
        ("Backend API", "Express, JWT, mysql2, Zod", "REST /api, validation, orchestration AI proxy", "Mount /api; helmet, CORS, middleware lỗi tập trung"),
        ("Admin web", "Express templates HTML", "Dashboard, CRUD người dùng/địa điểm, giám sát RAG", "Bảo vệ Basic Auth khi cấu hình ADMIN_BASIC_*"),
        ("RAG service", "FastAPI, pipeline retrieval", "Chat RAG, itinerary AI HTTP", "AI_RUNTIME_MODE/GEMINI theo env; khóa nội bộ tùy chọn"),
        ("MySQL", "InnoDB / relational", "Lưu trữ nghiệp vụ và metadata", "Pool kết nối trong backend/nodejs/src/db.js"),
    ]
    for r in rows:
        cells = tbl.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = v

    # ========== 3.2 ==========
    doc.add_heading("3.2. Thiết kế kiến trúc hệ thống", level=1)

    doc.add_heading("3.2.1. Mô hình kiến trúc ba lớp", level=2)
    add_para(
        doc,
        "Hệ thống được mô tả theo mô hình ba lớp kinh điển: (1) Lớp trình bày — giao diện Android "
        "(Activity/Fragment, ViewBinding, Navigation Component); (2) Lớp nghiệp vụ — backend Express "
        "xử lý quy tắc nghiệp vụ, phân quyền, định dạng DTO; (3) Lớp dữ liệu — MySQL và các repository "
        "truy vấn tham số hóa. Phân hệ RAG đóng vai trò “dịch vụ nghiệp vụ mở rộng” được backend gọi "
        "qua HTTP, không thay thế lớp persistence chính của ứng dụng.",
    )

    doc.add_heading("3.2.2. Sơ đồ kiến trúc hệ thống UML ba lớp", level=2)
    add_para(
        doc,
        "Ở góc nhìn triển khai-logic, có thể minh họa ba lớp như sau: Presentation (Android) chỉ "
        "phụ thuộc contract REST; Application/Business (Node) điều phối Auth, Destinations, "
        "Itineraries, Reviews và AI proxy; Data Access (Repositories + MySQL) cô lập SQL; Integration "
        "(RAG Client trong Node, WeatherService trên Android) là các adapter ra hệ thống ngoài.",
    )
    add_para(
        doc,
        "(Hình 3.2 — Kiến trúc client–server và AI/RAG: UNU_Trip_Kien_truc_Client_Server_AI_RAG.png.)",
        italic=True,
    )

    doc.add_heading("3.2.3. Mô tả luồng xử lý giữa Android – Backend – Database – RAG", level=2)
    flow = [
        "Người dùng đăng nhập: Android gửi POST /api/auth/login → backend xác thực, sinh JWT → client lưu token trong EncryptedSharedPreferences.",
        "Thao tác CRUD có xác thực: Retrofit gửi Authorization: Bearer … → middleware JWT → service → repository → MySQL.",
        "Luồng chatbot: Android POST /api/ai/rag-chat → Node forward FastAPI /rag/chat/simple (có retry/timeout) → kết quả trả JSON → ViewModel có thể fallback /api/ai/chat khi RAG lỗi.",
        "Luồng gợi ý lịch trình AI: Android gọi các endpoint /api/ai/* và /api/itineraries/* (preview, options, save-ai, create-from-selection) → Node có thể truy vấn catalog địa điểm (app_places) và/hoặc gọi RAG để sinh/cấu trúc nội dung.",
        "Thời tiết hiển thị tại chi tiết địa điểm: Android gọi trực tiếp Open-Meteo, không đi qua backend.",
    ]
    for f in flow:
        add_bullet(doc, f)

    add_para(
        doc,
        "(Hình 3.3 — Luồng chatbot RAG: UNU_Trip_Flow_Chatbot_RAG.png; luồng itinerary: UNU_Trip_Flow_AI_Itinerary_RAG.png.)",
        italic=True,
    )

    # ========== 3.3 ==========
    doc.add_heading("3.3. Thiết kế cơ sở dữ liệu mức logic", level=1)

    doc.add_heading("3.3.1. Nguyên tắc thiết kế dữ liệu", level=2)
    add_para(
        doc,
        "Dữ liệu được chuẩn hóa theo mô hình quan hệ, khóa chính–khóa ngoại được thể hiện rõ trong "
        "các câu truy vấn repository; tránh nhập SQL động không tham số hóa; tách bảng địa điểm hiển "
        "thị ứng dụng (app_places) khỏi kho tri thức phục vụ RAG (rag_knowledge_base) và bảng ánh xạ "
        "(place_id_map) để đồng bộ định danh giữa pipeline AI và dữ liệu app.",
    )

    doc.add_heading("3.3.2. Danh sách bảng dữ liệu chính", level=2)
    dt = doc.add_table(rows=1, cols=3)
    h = dt.rows[0].cells
    h[0].text = "Bảng"
    h[1].text = "Chức năng logic"
    h[2].text = "Liên quan ứng dụng"
    drows = [
        ("users", "Tài khoản, hồ sơ cơ bản, preferences_json", "Đăng nhập/đăng ký, profile"),
        ("app_places", "Địa điểm du lịch hiển thị trong app", "Danh sách, chi tiết, nearby/featured"),
        ("favorites", "Liên kết user ↔ địa điểm yêu thích", "Tab yêu thích"),
        ("reviews", "Đánh giá sao, nhận xét, ảnh đính kèm", "Chi tiết địa điểm"),
        ("itineraries", "Header lịch trình", "Danh sách/tạo lịch"),
        ("itinerary_days", "Ngày trong lịch", "Chi tiết lịch"),
        ("itinerary_items", "Điểm dừng trong ngày", "Chi tiết lịch"),
        ("place_images", "Ảnh địa điểm V2", "Giao diện media"),
        ("place_id_map", "Ánh xạ ID RAG/raw ↔ app_places", "Luồng AI đồng bộ"),
        ("rag_knowledge_base", "Corpus mô tả phục vụ ingest RAG", "Export corpus / pipeline"),
    ]
    for row in drows:
        c = dt.add_row().cells
        for i, v in enumerate(row):
            c[i].text = v

    doc.add_heading("3.3.3. Ràng buộc và quan hệ dữ liệu", level=2)
    add_para(
        doc,
        "Quan hệ logic: favorites và reviews tham chiếu users và destination_id trùng khóa app_places.id; "
        "itinerary_days khóa ngoại itineraries.id; itinerary_items khóa ngoại itinerary_days và "
        "destination_id → app_places. Các file migration trong database/migrations (001–004) định nghĩa "
        "cấu trúc app_places, rag_knowledge_base, place_images, place_id_map — nên đính kèm ERD minh họa.",
    )
    add_para(doc, "(Hình 3.4 — ERD rút gọn: UNU_Trip_ERD_rut_gon.png.)", italic=True)

    # ========== 3.4 ==========
    doc.add_heading("3.4. Thiết kế API hệ thống", level=1)

    doc.add_heading("3.4.1. Nguyên tắc thiết kế REST API", level=2)
    principles = [
        "Tài nguyên được đặt tên theo danh từ số nhiều (users, destinations, itineraries, reviews).",
        "Đăng nhập/đăng ký và health là các endpoint công khai; phần lớn /api/* yêu cầu JWT.",
        "Phản hồi JSON có trường success/message/data phù hợp từng module; lỗi HTTP có middleware thống nhất.",
        "Header trace: client có thể gửi X-Request-ID, X-Client-Version qua interceptor Retrofit để hỗ trợ log.",
        "Upload multimedia (avatar, ảnh review) dùng multipart; các luồng AI dùng JSON body có validate phía server.",
    ]
    for p in principles:
        add_bullet(doc, p)

    doc.add_heading("3.4.2. Danh sách API chính", level=2)
    at = doc.add_table(rows=1, cols=4)
    ah = at.rows[0].cells
    ah[0].text = "Nhóm"
    ah[1].text = "Phương thức"
    ah[2].text = "Đường dẫn (prefix /api)"
    ah[3].text = "Mục đích"
    apis = [
        ("Health", "GET", "/health, /health/ready", "Kiểm tra sống/sẵn sàng"),
        ("Auth", "POST", "/auth/register, /auth/login", "Đăng ký/đăng nhập JWT"),
        ("Auth", "POST", "/auth/logout", "Đăng xuất stateless (client bỏ token)"),
        ("Users", "GET/PUT/POST", "/users/profile, /users/stats, /users/preferences, /users/avatar", "Hồ sơ & avatar"),
        ("Favorites", "GET/POST/DELETE", "/users/favorites, /users/favorites/:id", "Yêu thích"),
        ("Destinations", "GET", "/destinations, /featured, /nearby, /:id", "Danh sách & chi tiết địa điểm"),
        ("Reviews", "GET/POST", "/destinations/:id/reviews, /reviews", "Đánh giá (+ảnh)"),
        ("Itineraries", "REST", "/itineraries, /:id, /:id/items", "CRUD lịch trình"),
        ("AI & Itinerary", "POST", "/ai/suggest-itinerary, /ai/rag-chat, /ai/chat, /ai/itinerary-preview, /ai/itinerary-options", "AI proxy"),
        ("AI & Itinerary", "POST", "/itineraries/save-ai, /create-from-selection, /create-from-option", "Tạo/ghép lịch AI"),
    ]
    for row in apis:
        c = at.add_row().cells
        for i, v in enumerate(row):
            c[i].text = v

    doc.add_heading("3.4.3. Thiết kế request/response cho API tiêu biểu", level=2)

    doc.add_heading("a) Đăng nhập", level=3)
    add_para(doc, "Request JSON: { \"email\": string, \"password\": string } (validate Zod).")
    add_para(doc, "Response thành công: { success, message, token, user } — user là DTO đã loại trường nhạy cảm.")
    add_para(doc, "Lỗi: 400 payload không hợp lệ; 401 sai thông tin đăng nhập.")

    doc.add_heading("b) Danh sách địa điểm", level=3)
    add_para(doc, "Query: page, limit, category, province, search, sort.")
    add_para(doc, "Response theo DestinationResponse: success, data[], total, page, limit — joined trạng thái is_favorite khi có user.")

    doc.add_heading("c) Tạo lịch trình", level=3)
    add_para(doc, "Request CreateItineraryRequest: title, description, startDate, endDate (khớp Models.kt và controller backend).")
    add_para(doc, "Response: ApiResponse<Itinerary> sau khi insert itineraries và logic ngày liên quan.")

    # ========== 3.5 ==========
    doc.add_heading("3.5. Thiết kế phân hệ AI/RAG", level=1)

    doc.add_heading("3.5.1. Mục tiêu phân hệ AI/RAG", level=2)
    add_para(
        doc,
        "Phân hệ nhằm cung cấp trả lời có căn cứ từ kho tri thức địa điểm, hỗ trợ chatbot và các luồng "
        "gợi ý/chỉnh sửa lịch trình. Yêu cầu phi chức năng bao gồm kiểm soát timeout/retry phía Node, "
        "rate limit phía FastAPI, và khả năng vận hành ở chế độ mock khi chưa bật Gemini.",
    )

    doc.add_heading("3.5.2. Kiến trúc phân hệ AI/RAG", level=2)
    add_para(
        doc,
        "FastAPI app “UnuTrip RAG” expose các router health, rag, admin; middleware InternalApiKeyMiddleware "
        "và RagRateLimitMiddleware. Node.js dùng ragUpstream (retry, timeout RAG_FETCH_TIMEOUT_MS) để gọi các "
        "path như /rag/chat/simple. Cấu hình ENABLE_GEMINI, GEMINI_API_KEY, AI_RUNTIME_MODE đọc từ biến môi "
        "trường — mặc định có thể là mock; báo cáo nên nêu rõ phụ thuộc cấu hình triển khai.",
    )

    doc.add_heading("3.5.3. Luồng xử lý gợi ý lịch trình", level=2)
    add_para(
        doc,
        "Android (AISuggestFragment / các fragment AI tour) gọi POST /api/ai/suggest-itinerary hoặc các endpoint "
        "preview/options tương ứng → Node build prompt, có thể gọi model cục bộ qua AI_MODEL_URL hoặc fallback "
        "RAG JSON → parse và trả itinerary có cấu trúc → client có thể POST /api/itineraries/save-ai để lưu.",
    )

    doc.add_heading("3.5.4. Thiết kế dữ liệu đầu vào và đầu ra", level=2)
    add_para(
        doc,
        "Đầu vào thường là ngữ cảnh người dùng (độ dài chuyến, sở thích, điểm khởi hành…) và/hoặc tập địa điểm "
        "catalog trích từ app_places qua repository AI. Đầu ra là JSON itinerary (ngày, thứ tự, destination_id, "
        "khung giờ gợi ý) được kiểm tra trước khi map vào bảng itineraries/itinerary_days/itinerary_items.",
    )

    doc.add_heading("3.5.5. Kiểm soát lỗi và chuẩn hóa kết quả AI", level=2)
    controls = [
        "HTTP exception handler FastAPI trả success=false kèm request_id.",
        "Node: retry transient khi gọi RAG; log lỗi downstream.",
        "Chatbot: fallback sang /api/ai/chat khi RAG lỗi; ViewModel có chuỗi validate/repair phía client-service.",
        "Itinerary JSON không hợp lệ → phản hồi 500 có thông điệp “AI trả về dữ liệu không hợp lệ” (theo controller hiện có).",
    ]
    for c in controls:
        add_bullet(doc, c)

    # ========== 3.6 ==========
    doc.add_heading("3.6. Thiết kế giao diện ứng dụng", level=1)

    doc.add_heading("3.6.1. Nguyên tắc thiết kế giao diện", level=2)
    add_para(
        doc,
        "Giao diện Android chủ đạo dùng Fragment + XML + ViewBinding; Navigation Component thống nhất luồng "
        "giữa các màn hình; BottomNavigation cho các tab chính (Trang chủ, Địa điểm, Lịch trình, Chatbot, Hồ sơ). "
        "Typography và theme Material được áp dụng nhất quán; ảnh địa điểm và card layout tuân card/list pattern.",
    )

    subs = [
        ("3.6.2", "Mockup màn hình đăng nhập/đăng ký", "AuthActivity — form đơn giản, chuyển MainActivity sau khi lưu phiên."),
        ("3.6.3", "Mockup màn hình trang chủ", "HomeFragment — featured, nearby, ô tìm kiếm."),
        ("3.6.4", "Mockup màn hình danh sách/chi tiết địa điểm", "DestinationListFragment, DestinationDetailFragment — ảnh, tag, review, yêu thích, thêm vào lịch."),
        ("3.6.5", "Mockup màn hình yêu thích", "DestinationListFragment với tham số chỉ hiển thị favorites."),
        ("3.6.6", "Mockup màn hình AI gợi ý lịch trình", "AIItineraryRequest/Options/Editor và AISuggestFragment."),
        ("3.6.7", "Mockup màn hình lịch trình cá nhân", "ItineraryFragment + ItineraryDetailFragment."),
        ("3.6.8", "Mockup màn hình hồ sơ người dùng", "ProfileFragment — avatar, stats, điều hướng settings."),
    ]
    for num, title, desc in subs:
        doc.add_heading(f"{num}. {title}", level=2)
        add_para(doc, f"Mô tả chức năng theo mã nguồn: {desc}")
        add_para(doc, "(Chèn ảnh chụp màn hình thực tế hoặc thiết kế Figma tương ứng.)", italic=True)

    add_para(
        doc,
        "Tham khảo bố cục bản đồ trong repo: UNU_Trip_Thiet_ke_man_hinh_Ban_do.png; chatbot: UNU_Trip_Thiet_ke_Chatbot_AI.png.",
        italic=True,
    )

    # ========== 3.7 ==========
    doc.add_heading("3.7. Thiết kế luồng giao diện người dùng", level=1)

    doc.add_heading("3.7.1. Luồng đăng nhập/đăng ký", level=2)
    add_para(doc, "Launcher AuthActivity → POST auth/login|register → EncryptedSharedPreferences lưu JWT → start MainActivity.")

    doc.add_heading("3.7.2. Luồng tìm kiếm và xem địa điểm", level=2)
    add_para(doc, "Home/List → GET destinations → Detail → MapFragment (OSMDroid, GPS FusedLocation) hoặc Intent Google Maps khi chỉ đường ngoài app.")

    doc.add_heading("3.7.3. Luồng tạo lịch trình bằng AI", level=2)
    add_para(doc, "Itinerary tab → chọn flow AI (request/options/editor hoặc AISuggest) → các POST /api/ai/* và /api/itineraries/* → refresh danh sách.")

    doc.add_heading("3.7.4. Luồng quản lý yêu thích/lịch trình", level=2)
    add_para(doc, "Favorite toggle qua POST/DELETE users/favorites; chọn lịch trong dialog và thêm điểm qua POST itineraries/:id/items.")
    add_para(doc, "(Hình 3.5 — Điều hướng Android: UNU_Trip_Android_Navigation.png.)", italic=True)

    # ========== 3.8 ==========
    doc.add_heading("3.8. Thiết kế bảo mật hệ thống", level=1)

    doc.add_heading("3.8.1. Xác thực và phân quyền", level=2)
    add_para(doc, "Người dùng cuối: JWT ký bằng JWT_SECRET; middleware auth cho các route nhạy cảm. Admin: Basic Auth khi cấu hình đủ ADMIN_BASIC_USER/PASS.")

    doc.add_heading("3.8.2. Bảo vệ API", level=2)
    add_para(doc, "Helmet, CSP nonce cho admin; CORS mặc định (khuyến nghị thu hẹp production). RAG: Internal API Key qua header khi bật.")

    doc.add_heading("3.8.3. Kiểm tra dữ liệu đầu vào", level=2)
    add_para(doc, "Validate schema (Zod) ở controller; SQL parameterized qua mysql2; multipart giới hạn kích thước theo upload middleware.")

    doc.add_heading("3.8.4. Bảo vệ dữ liệu người dùng", level=2)
    add_para(doc, "Mật khẩu bcrypt; token lưu EncryptedSharedPreferences; prod flavor tắt cleartext HTTP — cần HTTPS hoặc network security phù hợp.")

    # ========== 3.9 ==========
    doc.add_heading("3.9. Thiết kế triển khai hệ thống", level=1)

    doc.add_heading("3.9.1. Mô hình triển khai", level=2)
    add_para(doc, "Docker Compose có thể gói MySQL, Redis, backend Node, RAG; Android emulator trỏ 10.0.2.2:3000/api; thiết bị thật dùng IP LAN trong local.properties.")

    doc.add_heading("3.9.2. Thành phần triển khai", level=2)
    add_bullet(doc, "Container/service: mysql:3306, redis:6379, backend:3000, rag:8001 (publish ports theo .env).")
    add_bullet(doc, "Artifact Android: flavor devDebug/prodRelease — BASE_URL và cờ cleartext khác nhau.")

    doc.add_heading("3.9.3. Luồng giao tiếp khi vận hành", level=2)
    add_para(doc, "Client HTTPS/HTTP JSON ↔ Node /api ↔ MySQL; Node ↔ RAG nội bộ qua RAG_BASE_URL; Admin browser ↔ /admin HTML ↔ MySQL/RAG admin endpoints.")
    add_para(doc, "(Hình 3.6 — Triển khai dev local: UNU_Trip_Trien_khai_Dev_Local.png.)", italic=True)

    # ========== 3.10 ==========
    doc.add_heading("3.10. Kết luận chương", level=1)
    add_para(
        doc,
        "Chương III đã trình bày thiết kế tổng thể và chi tiết các lớp của hệ thống UNU Trip: kiến trúc ba lớp "
        "với tích hợp RAG độc lập; mô hình dữ liệu quan hệ và các bảng cốt lõi phục vụ mobile API; hợp đồng REST "
        "nhất quán giữa Android và backend; phân hệ AI/RAG với cơ chế fallback và kiểm soát lỗi; nguyên tắc giao diện "
        "và luồng điều hướng; các biện pháp bảo mật cơ bản; cùng phác thảo triển khai Docker và biến môi trường. "
        "Chương tiếp theo sẽ đi sâu hiện thực hóa, kiểm thử và đánh giá kết quả trên nền các thiết kế này.",
    )

    try:
        doc.save(OUT_PATH)
        print("Saved:", OUT_PATH)
    except PermissionError:
        doc.save(ALT_PATH)
        print(
            "Không ghi được file gốc (đang mở trong Word?). Đã lưu:",
            ALT_PATH,
            "\nĐóng Word và chạy lại script hoặc đổi tên ALT_PATH -> OUT_PATH.",
        )


if __name__ == "__main__":
    build()
